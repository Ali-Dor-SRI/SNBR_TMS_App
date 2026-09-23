"""Parse QtracP per-stimulus Excel exports into plain dicts.

QtracP can write every pulse of a recording to an ``.xlsx`` workbook: one sheet
per recorded variable (``T`` test stimulus in %MSO, ``P`` peak MEP in mV, ``L``
latency in ms, ``D`` conditioning-test delay in ms, which is the ISI), and inside
each sheet one ``(elapsed time, value)`` column pair per Qtrac channel under a
``Chan  N`` header. The workbook is matched to its ``.MEM`` by the acquisition
token that ends both file stems (``TP3C60922B``), never by the participant
number, which is absent from some names and wrong in others.

What comes out is the per-pulse detail behind the profile values the ``.MEM``
already holds, plus a recomputation of those values by Qtrac's own rules, so a
mismatch is visible:

* **Threshold-tracked measures** (T-SICI, T-SICF). Each ISI channel's value is
  Qtrac's "Log regression": stimulus regressed on ln(MEP) with weights
  ``1 - |log10(MEP / 0.2 mV)|`` (zero outside one tenth to ten times the target
  and outside 0.01 to 2 mV), read at 0.2 mV, then expressed as a percentage of
  the RMT200 tracked in parallel on channel 1 during the same phase, **not** the
  header RMT200. Each pulse's reading is its test stimulus on the same scale.
* **Amplitude measures** (A-SICI, A-SICF). Each ISI channel's value is the
  geometric mean of its MEPs as a percentage of the test-alone baseline: the
  geometric mean of the channel-5 MEPs delivered during the phase, counting from
  the one just before the first pair. Each pulse's reading is its MEP on the
  same scale.

Channels 20, 21 and 22 hold traces Qtrac rejected (spontaneous EMG, a long
stimulus artefact, a missing stimulus) and are excluded, as they are in the MEM.
The CSP exports in the same folders are waveform workbooks (``W...`` sheets) and
are recognised and skipped rather than parsed.

Parsers return plain dicts, never DataFrames.
"""

from __future__ import annotations

import math
import re
from collections import defaultdict
from pathlib import Path

import numpy as np

from parser.mem_parser import iter_files

# The acquisition token that ends every Qtrac file stem: two to four letters, a
# digit, a letter, five digits, a letter (``TP3C60922B``, ``TG2C40422A``).
_TOKEN_RE = re.compile(r"([A-Z]{2,4}\d[A-Z]\d{5}[A-Z])$")
_CHAN_RE = re.compile(r"^Chan\s+(\d+)$")

TARGET_MV = 0.2                     # RMT200 target response
CONDITIONING_FRACTION = 0.7         # conditioning pulse as a fraction of the tracked RMT200
TEST_ALONE_CHAN = 1                 # RMT200 tracking / test-alone pulse of the threshold phases
AMPLITUDE_TEST_CHAN = 5             # RMT1000 tracking / test-alone pulse of the amplitude phases
CONDITIONED_CHANS = (4, 6, 7, 8, 9, 11, 12, 13, 14, 15, 16, 17, 18, 19)
SKIP_CHANS = {20: "emg", 21: "artefact", 22: "missing"}
CSP_DELAY = -300.0                  # the CSP staircase is exported with this delay

KIND_STIMULI = "stimuli"            # D/P/L/T: thresholds and amplitudes
KIND_STIMULI_NO_T = "stimuli_no_t"  # D/P plus a waveform sheet: amplitudes only
KIND_WAVEFORMS = "waveforms"        # the CSP exports (Wxxxx ... W sheets)
KIND_UNKNOWN = "unknown"

PULSE_MEASURES = ("t_sici", "t_sicf", "a_sici", "a_sicf")
THRESHOLD_MEASURES = ("t_sici", "t_sicf")
AMPLITUDE_MEASURES = ("a_sici", "a_sicf")


# ---------------------------------------------------------------------------
# Discovery and matching
# ---------------------------------------------------------------------------

def acquisition_token(name: str | Path) -> str | None:
    """The Qtrac acquisition token ending *name*'s stem, upper-cased, or None."""
    match = _TOKEN_RE.search(Path(str(name)).stem.strip().upper())
    return match.group(1) if match else None


def tokens_in(source_file_cell) -> set[str]:
    """Tokens of every file named in a ``source_file`` cell (``"; "``-joined)."""
    if source_file_cell is None:
        return set()
    if isinstance(source_file_cell, float) and math.isnan(source_file_cell):
        return set()
    found = set()
    for name in str(source_file_cell).split(";"):
        token = acquisition_token(name.strip())
        if token:
            found.add(token)
    return found


def iter_xlsx_files(input_dir, recursive: bool = False) -> list[Path]:
    """Every ``*.xlsx`` under the folder(s), de-duplicated, missing folders skipped."""
    return iter_files(input_dir, "*.xlsx", recursive)


def workbook_kind(path: str | Path) -> str:
    """Classify a workbook by its sheets without reading its data.

    Only the sheet names and each sheet's first row are read, so a folder of
    sixty workbooks classifies in well under a second.
    """
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return KIND_UNKNOWN
    try:
        names = list(wb.sheetnames)
        stimulus_sheets = set()
        for name in names:
            if name in ("D", "P", "L", "T"):
                ws = wb[name]
                first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
                if any(isinstance(h, str) and _CHAN_RE.match(h) for h in first):
                    stimulus_sheets.add(name)
    finally:
        wb.close()
    if {"D", "P", "T"} <= stimulus_sheets:
        return KIND_STIMULI
    if {"D", "P"} <= stimulus_sheets:
        return KIND_STIMULI_NO_T
    if names and all(n == "QPP" or n.startswith("W") for n in names):
        return KIND_WAVEFORMS
    return KIND_UNKNOWN


def index_workbooks(input_dir, recursive: bool = False) -> dict[str, Path]:
    """``{token: workbook}`` for the stimulus exports under the folder(s).

    Waveform (CSP) workbooks and files without a token are left out. When one
    recording was exported more than once, the export that carries the
    stimulus column wins, then the newest file.
    """
    ranked: dict[str, tuple[int, float, Path]] = {}
    for path in iter_xlsx_files(input_dir, recursive):
        token = acquisition_token(path.name)
        if not token:
            continue
        kind = workbook_kind(path)
        if kind == KIND_STIMULI:
            rank = 2
        elif kind == KIND_STIMULI_NO_T:
            rank = 1
        else:
            continue
        try:
            mtime = path.stat().st_mtime
        except OSError:
            mtime = 0.0
        current = ranked.get(token)
        if current is None or (rank, mtime) > current[:2]:
            ranked[token] = (rank, mtime, path)
    return {token: entry[2] for token, entry in ranked.items()}


def workbook_for_source_files(index: dict[str, Path], source_files) -> Path | None:
    """The workbook matching any file named in *source_files* (names or cells)."""
    if isinstance(source_files, str):
        source_files = [source_files]
    for cell in source_files or []:
        for token in sorted(tokens_in(cell)):
            if token in index:
                return index[token]
    return None


# ---------------------------------------------------------------------------
# Reading the workbook
# ---------------------------------------------------------------------------

def read_stimulus_workbook(path: str | Path) -> dict[int, list[dict]]:
    """``{channel: [pulse, ...]}`` with ``time`` (min), ``T``, ``P``, ``L``, ``D`` per pulse.

    Sheets that were not exported are absent from the pulse dicts (a workbook
    of the ``D``/``P``/waveform variant has no ``T`` or ``L``).
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {
            name: list(wb[name].iter_rows(values_only=True))
            for name in wb.sheetnames if name in ("D", "P", "L", "T")
        }
    finally:
        wb.close()
    value_sheets = [
        name for name, rows in sheets.items()
        if rows and any(isinstance(h, str) and _CHAN_RE.match(h) for h in rows[0])
    ]
    base_name = next((n for n in ("T", "P", "D") if n in value_sheets), None)
    if base_name is None:
        return {}
    base = sheets[base_name]
    channels: dict[int, list[dict]] = {}
    for col, header in enumerate(base[0]):
        match = _CHAN_RE.match(header) if isinstance(header, str) else None
        if not match or col == 0:
            continue
        chan = int(match.group(1))
        pulses = []
        for i in range(1, len(base)):
            row = base[i]
            if col >= len(row) or row[col - 1] is None:
                continue
            pulse = {"time": float(row[col - 1])}
            for name in value_sheets:
                sheet_row = sheets[name][i] if i < len(sheets[name]) else ()
                value = sheet_row[col] if col < len(sheet_row) else None
                pulse[name] = float(value) if isinstance(value, (int, float)) else None
            pulses.append(pulse)
        if pulses:
            channels[chan] = sorted(pulses, key=lambda p: p["time"])
    return channels


# ---------------------------------------------------------------------------
# Qtrac's estimators
# ---------------------------------------------------------------------------

def qtrac_log_regression(stimuli, responses, target: float = TARGET_MV) -> tuple[float, float]:
    """Qtrac's "Log regression" threshold: ``(threshold, standard error)``.

    Stimulus regressed on ln(response) so that with no correlation the estimate
    is the (weighted) mean stimulus; weights ``1 - |log10(response / target)|``,
    clipped at zero, and zero for responses outside 0.01 to 2 mV. NaN when fewer
    than two pulses carry weight.
    """
    T = np.asarray(stimuli, dtype=float)
    P = np.asarray(responses, dtype=float)
    safe = np.clip(P, 1e-9, None)
    weights = np.clip(1.0 - np.abs(np.log10(safe / target)), 0.0, None)
    weights[(P < 0.01) | (P > 2.0) | ~np.isfinite(P) | ~np.isfinite(T)] = 0.0
    used = int((weights > 0).sum())
    if used < 2:
        return float("nan"), float("nan")
    x = np.log(safe)
    total = weights.sum()
    mean_x = (weights * x).sum() / total
    mean_t = (weights * T).sum() / total
    sxx = (weights * (x - mean_x) ** 2).sum()
    slope = (weights * (x - mean_x) * (T - mean_t)).sum() / sxx if sxx > 0 else 0.0
    threshold = mean_t + slope * (math.log(target) - mean_x)
    residuals = T - (mean_t + slope * (x - mean_x))
    dof = max(total * (used - 2) / used, 1e-9)
    variance = (weights * residuals ** 2).sum() / dof
    se = (
        math.sqrt(variance * (1.0 / total + (math.log(target) - mean_x) ** 2 / sxx))
        if sxx > 0 else float("nan")
    )
    return float(threshold), float(se)


def geometric_mean(values) -> float:
    """Geometric mean of the positive, finite values (NaN when there are none)."""
    arr = np.asarray([v for v in values if v is not None], dtype=float)
    arr = arr[np.isfinite(arr) & (arr > 0)]
    return float(np.exp(np.log(arr).mean())) if len(arr) else float("nan")


# ---------------------------------------------------------------------------
# Phases
# ---------------------------------------------------------------------------

def _spans(pulses, gap_minutes: float = 3.0) -> list[tuple[float, float]]:
    out: list[tuple[float, float]] = []
    start = end = None
    for pulse in sorted(pulses, key=lambda p: p["time"]):
        t = pulse["time"]
        if start is None:
            start = end = t
        elif t - end > gap_minutes:
            out.append((start, end))
            start = end = t
        else:
            end = t
    if start is not None:
        out.append((start, end))
    return out


def _in_spans(t: float, spans, pad: float = 0.5) -> bool:
    return any(s - pad <= t <= e + pad for s, e in spans)


def split_phases(channels: dict[int, list[dict]]) -> dict[str, dict[float, list[dict]]]:
    """Sort the conditioned pulses into the four measures, keyed by ISI.

    A conditioned pulse delivered while channel 5 is running belongs to an
    amplitude phase (its test stimulus is held fixed); any other belongs to a
    threshold-tracking phase. The sign of the delay tells SICI from SICF.
    """
    amplitude_spans = _spans(channels.get(AMPLITUDE_TEST_CHAN, []))
    phases: dict[str, dict[float, list[dict]]] = {m: defaultdict(list) for m in PULSE_MEASURES}
    for chan in CONDITIONED_CHANS:
        for pulse in channels.get(chan, []):
            delay = pulse.get("D")
            if delay is None or delay == 0 or delay <= CSP_DELAY + 1:
                continue
            amplitude = _in_spans(pulse["time"], amplitude_spans)
            if delay > 0:
                measure = "a_sici" if amplitude else "t_sici"
            else:
                measure = "a_sicf" if amplitude else "t_sicf"
            phases[measure][round(abs(delay), 2)].append(pulse)
    return {
        measure: {isi: sorted(v, key=lambda p: p["time"]) for isi, v in by_isi.items()}
        for measure, by_isi in phases.items()
    }


def _phase_span(by_isi) -> tuple[float, float]:
    times = [p["time"] for pulses in by_isi.values() for p in pulses]
    return min(times), max(times)


def _reference_pulses(channels, chan: int, span) -> list[dict]:
    """The test-alone pulses of *chan* during *span*, plus the one just before it."""
    pulses = sorted(channels.get(chan, []), key=lambda p: p["time"])
    t0, t1 = span
    before = [p for p in pulses if p["time"] < t0]
    return ([before[-1]] if before else []) + [p for p in pulses if t0 <= p["time"] <= t1]


def _latest_stimulus(pulses, t: float):
    prior = [p for p in pulses if p["time"] <= t and p.get("T") is not None and p["T"] > 5]
    return prior[-1]["T"] if prior else None


# ---------------------------------------------------------------------------
# The readings
# ---------------------------------------------------------------------------

def _threshold_block(channels, by_isi, warnings: list[str], label: str) -> dict | None:
    if any(p.get("T") is None for pulses in by_isi.values() for p in pulses):
        warnings.append(f"{label}: the workbook has no test-stimulus (T) sheet")
        return None
    span = _phase_span(by_isi)
    reference_pulses = [
        p for p in _reference_pulses(channels, TEST_ALONE_CHAN, span)
        if p.get("T") is not None and p.get("P") is not None
    ]
    reference, _ = qtrac_log_regression(
        [p["T"] for p in reference_pulses], [p["P"] for p in reference_pulses],
    )
    if not math.isfinite(reference) or reference <= 0:
        warnings.append(f"{label}: no channel-1 pulses to reference the thresholds against")
        return None

    # Two recordings on the share export the paired-channel T as the increment
    # above the conditioning pulse rather than the test pulse itself; the pairs
    # then sit far below the RMT. Restore the test pulse before regressing.
    paired = [p["T"] for pulses in by_isi.values() for p in pulses]
    increment_mode = bool(paired) and float(np.median(paired)) < 0.6 * reference
    if increment_mode:
        warnings.append(
            f"{label}: paired-pulse stimuli were exported as increments above the "
            f"conditioning pulse; {CONDITIONING_FRACTION:.0%} of the tracked RMT200 was added back"
        )
    test_alone = channels.get(TEST_ALONE_CHAN, [])

    def stimulus_of(pulse) -> float:
        value = float(pulse["T"])
        if increment_mode:
            latest = _latest_stimulus(test_alone, pulse["time"])
            if latest is not None:
                value += round(CONDITIONING_FRACTION * latest)
        return value

    isis: dict[float, dict] = {}
    for isi, pulses in sorted(by_isi.items()):
        stimuli = [stimulus_of(p) for p in pulses]
        responses = [p["P"] if p.get("P") is not None else float("nan") for p in pulses]
        threshold, se = qtrac_log_regression(stimuli, responses)
        isis[isi] = {
            "value": threshold / reference * 100.0 if math.isfinite(threshold) else None,
            "se": se / reference * 100.0 if math.isfinite(se) else None,
            "pulses": [
                {
                    "y": stimulus / reference * 100.0,
                    "above_target": response >= TARGET_MV if math.isfinite(response) else False,
                    "time": p["time"],
                    "stimulus": stimulus,
                    "mep": response,
                }
                for p, stimulus, response in zip(pulses, stimuli, responses)
            ],
        }
    return {
        "reference": float(reference),
        "reference_n": len(reference_pulses),
        "reference_kind": "parallel_rmt200",
        "increment_mode": increment_mode,
        "isis": isis,
    }


def _amplitude_block(channels, by_isi, warnings: list[str], label: str) -> dict | None:
    span = _phase_span(by_isi)
    baseline_pulses = [
        p for p in _reference_pulses(channels, AMPLITUDE_TEST_CHAN, span)
        if p.get("P") is not None
    ]
    baseline = geometric_mean([p["P"] for p in baseline_pulses])
    if not math.isfinite(baseline) or baseline <= 0:
        warnings.append(f"{label}: no channel-5 test-alone pulses to use as the baseline")
        return None
    isis: dict[float, dict] = {}
    for isi, pulses in sorted(by_isi.items()):
        meps = [p["P"] if p.get("P") is not None else float("nan") for p in pulses]
        isis[isi] = {
            "value": geometric_mean(meps) / baseline * 100.0,
            "se": None,
            "pulses": [
                {
                    "y": mep / baseline * 100.0 if math.isfinite(mep) else None,
                    "above_target": True,
                    "time": p["time"],
                    "stimulus": p.get("T"),
                    "mep": mep,
                }
                for p, mep in zip(pulses, meps)
            ],
        }
        isis[isi]["pulses"] = [pl for pl in isis[isi]["pulses"] if pl["y"] is not None]
    return {
        "reference": float(baseline),
        "reference_n": len(baseline_pulses),
        "reference_kind": "test_alone_baseline",
        "increment_mode": False,
        "isis": isis,
    }


def extract_pulse_readings(path: str | Path) -> dict:
    """Per-pulse readings and recomputed profile values for one workbook.

    Returns a plain dict::

        {
          "path": str, "token": str | None, "kind": str,
          "measures": {measure: block | None},   # see _threshold_block / _amplitude_block
          "skipped_traces": {"emg": n, "artefact": n, "missing": n},
          "warnings": [str, ...],
        }

    A measure the recording did not run is absent from ``measures``; one that
    ran but could not be recomputed is present as ``None`` with a warning.
    """
    path = Path(path)
    warnings: list[str] = []
    kind = workbook_kind(path)
    result = {
        "path": str(path),
        "token": acquisition_token(path.name),
        "kind": kind,
        "measures": {},
        "skipped_traces": {},
        "warnings": warnings,
    }
    if kind not in (KIND_STIMULI, KIND_STIMULI_NO_T):
        warnings.append("not a per-stimulus export (no D/P sheets with channel columns)")
        return result
    channels = read_stimulus_workbook(path)
    phases = split_phases(channels)
    for measure in PULSE_MEASURES:
        by_isi = phases[measure]
        if not by_isi:
            continue
        label = measure.upper().replace("_", "-")
        if measure in THRESHOLD_MEASURES:
            result["measures"][measure] = _threshold_block(channels, by_isi, warnings, label)
        else:
            result["measures"][measure] = _amplitude_block(channels, by_isi, warnings, label)
    result["skipped_traces"] = {
        name: len(channels.get(chan, [])) for chan, name in SKIP_CHANS.items()
    }
    return result


def readings_for_source_file(readings_by_token: dict[str, dict], source_file_cell) -> dict | None:
    """The readings dict whose token appears in a row's ``source_file`` cell."""
    for token in sorted(tokens_in(source_file_cell)):
        if token in readings_by_token:
            return readings_by_token[token]
    return None

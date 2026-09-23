"""Read the enrolment workbook into an in-memory roster. Read-only, always.

The workbook is the master list of who should receive a report. It is opened
with openpyxl in read-only mode, the header row is located by the configured
``Patient ID`` header, and every data row becomes an :class:`Identity` keyed
by its :class:`ParticipantKey`. Rows that cannot be keyed, and duplicate keys,
are returned as problems rather than dropped silently — a typo in one cell is
a fact the operator needs to see, not something to guess around.

Nothing here is ever written back to the workbook.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from dispatch.identity import Identity, ParticipantKey, parse_patient_id


class RosterError(Exception):
    """The workbook could not be read as a roster (missing sheet/column)."""


@dataclass(frozen=True)
class RosterProblem:
    row: int            # 1-based worksheet row
    patient_id_raw: str
    message: str


@dataclass
class Roster:
    entries: dict[ParticipantKey, Identity] = field(default_factory=dict)
    problems: list[RosterProblem] = field(default_factory=list)
    source: str = ""
    sheet: str = ""

    def __len__(self) -> int:
        return len(self.entries)

    def lookup(self, key: ParticipantKey) -> tuple[Identity | None, str]:
        """Find the identity for a report's key.

        Returns ``(identity, reason)``; *reason* is ``""`` on a match and
        otherwise says why there is none. A report with no study token is
        matched on number alone, but only when exactly one roster row has
        that number — participant numbers repeat between studies.
        """
        if key.study:
            for cand, ident in self.entries.items():
                if cand.number == key.number and cand.study.upper() == key.study.upper():
                    return ident, ""
            return None, f"{key.label()} is not in the roster"
        matches = [ident for cand, ident in self.entries.items() if cand.number == key.number]
        if len(matches) == 1:
            return matches[0], ""
        if not matches:
            return None, f"participant {key.number:03d} is not in the roster"
        studies = ", ".join(sorted(m.key.study for m in matches))
        return None, (
            f"participant {key.number:03d} matches several studies ({studies}) "
            f"and the report names none"
        )


def _norm(text) -> str:
    return " ".join(str(text or "").strip().lower().split())


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_roster(
    path: str | Path,
    sheet_name: str,
    col_patient_id: str,
    col_patient_name: str,
    col_mrn: str,
) -> Roster:
    """Read *sheet_name* of the workbook at *path* into a :class:`Roster`.

    Header matching is case- and whitespace-insensitive. The header row is the
    first row that contains the patient-ID header; the other two columns must
    be on that same row.
    """
    import openpyxl  # imported here so the package imports without it

    path = Path(path)
    if not path.is_file():
        raise RosterError(f"Roster workbook not found: {path}")

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        wanted = _norm(sheet_name)
        ws = None
        for name in wb.sheetnames:
            if _norm(name) == wanted:
                ws = wb[name]
                break
        if ws is None:
            raise RosterError(
                f"Sheet {sheet_name!r} not found in {path.name}. "
                f"Sheets present: {', '.join(wb.sheetnames)}"
            )

        targets = {
            "id": _norm(col_patient_id),
            "name": _norm(col_patient_name),
            "mrn": _norm(col_mrn),
        }
        columns: dict[str, int] = {}
        header_row = None
        rows = ws.iter_rows(values_only=True)
        for row_index, row in enumerate(rows, start=1):
            normed = [_norm(v) for v in row]
            if targets["id"] in normed:
                for role, header in targets.items():
                    if header in normed:
                        columns[role] = normed.index(header)
                header_row = row_index
                break
        if header_row is None:
            raise RosterError(
                f"No row in sheet {ws.title!r} contains a {col_patient_id!r} header"
            )
        missing = [
            {"name": col_patient_name, "mrn": col_mrn}[role]
            for role in ("name", "mrn") if role not in columns
        ]
        if missing:
            raise RosterError(
                f"Header row {header_row} of sheet {ws.title!r} lacks column(s): "
                f"{', '.join(repr(m) for m in missing)}"
            )

        roster = Roster(source=str(path), sheet=ws.title)
        for row_index, row in enumerate(rows, start=header_row + 1):
            cells = list(row)

            def col(role):
                idx = columns[role]
                return _cell_text(cells[idx]) if idx < len(cells) else ""

            raw_id, name, mrn = col("id"), col("name"), col("mrn")
            if not raw_id and not name and not mrn:
                continue
            key = parse_patient_id(raw_id)
            if key is None:
                roster.problems.append(RosterProblem(
                    row_index, raw_id,
                    "Patient ID does not end in a study token and number"
                    if raw_id else "Patient ID is blank",
                ))
                continue
            if key in roster.entries:
                roster.problems.append(RosterProblem(
                    row_index, raw_id,
                    f"duplicate of {key.label()} (row kept: first occurrence)",
                ))
                continue
            roster.entries[key] = Identity(
                key=key, patient_id_raw=raw_id, name=name, mrn=mrn,
            )
        return roster
    finally:
        wb.close()

"""A REDCap import row must name the record it is meant to update.

REDCap matches an import row to its record on the ``record_id`` **string**, and
this project's keys are zero-padded to three digits -- ``001``, ``005``,
``042``. A row keyed ``5`` therefore does not name record ``005``: REDCap
creates a *second* record and the visit's values land in it, leaving the real
one untouched. That is what these tests pin, in both directions:

* an **existing** record is named by REDCap's own key, copied verbatim out of
  the data export, whatever width or format that project uses;
* a **new** record -- a participant REDCap has never seen, which this import
  creates -- has no key yet, so the export decides it and writes it zero-padded
  to three digits, matching how the lab numbers participants in the ``.MEM``
  filenames, the report titles and the export names.

Keeping the key verbatim depends on reading it as text: ``record_id`` looks
numeric, so pandas types the column ``int64`` unless told otherwise and the
padding is gone before the diff even starts.

Only ``reports/redcap_exporter.py`` is affected. ``processing/redcap_mapper.py``
renames ``ID`` to ``record_id`` for new and existing rows alike and never reads
the REDCap export, so it cannot know either key.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from parser.mem_parser import TSICI_ISIS, initialize_record
from reports.redcap_exporter import (
    RECORD_ID_WIDTH,
    format_record_id,
    generate_redcap_import,
)

VISIT_DATE = "15/04/2026"
VISIT_DATE_ISO = "2026-04-15"


# ---------------------------------------------------------------------------
# The rule itself
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("pid,expected", [
    (1, "001"),
    (10, "010"),
    (100, "100"),
    (7, "007"),
    (80, "080"),
    (999, "999"),
])
def test_a_new_record_id_is_padded_to_three_digits(pid, expected):
    assert format_record_id(pid, is_new=True) == expected


@pytest.mark.parametrize("pid,key", [
    (5, "005"),
    (42, "042"),
    (158, "158"),
])
def test_an_existing_record_id_is_redcaps_own_key(pid, key):
    """The key has to arrive back exactly as REDCap holds it."""
    assert format_record_id(pid, is_new=False, existing_key=key) == key


@pytest.mark.parametrize("key", ["42", "0042", "00042"])
def test_an_existing_key_is_copied_not_reformatted(key):
    """A project numbering its records some other way keeps its own width."""
    assert format_record_id(42, is_new=False, existing_key=key) == key


def test_an_existing_key_is_never_the_bare_number_when_redcap_pads_it():
    """The defect itself: ``5`` does not name record ``005``, it creates one."""
    assert format_record_id(5, is_new=False, existing_key="005") != 5
    assert format_record_id(5, is_new=False, existing_key="005") != "5"


def test_surrounding_whitespace_in_the_key_is_dropped():
    assert format_record_id(5, is_new=False, existing_key=" 005 ") == "005"


def test_an_existing_record_with_no_key_falls_back_to_the_padded_form():
    """Never a bare integer: this project's keys are padded, so guess padded."""
    assert format_record_id(5, is_new=False) == "005"
    assert format_record_id(5, is_new=False, existing_key="") == "005"


@pytest.mark.parametrize("pid", [1, 5, 42, 100])
def test_padding_never_changes_which_participant_is_named(pid):
    assert int(format_record_id(pid, is_new=True)) == pid


def test_a_number_wider_than_the_pad_is_not_truncated():
    """Four digits stay four digits rather than losing the leading one."""
    assert format_record_id(1234, is_new=True) == "1234"


def test_the_pad_width_matches_the_rest_of_the_app():
    """The .MEM filenames and export names all use three digits."""
    assert RECORD_ID_WIDTH == 3
    assert format_record_id(5, is_new=True) == f"{5:0{RECORD_ID_WIDTH}d}"


# ---------------------------------------------------------------------------
# End to end, through the exporter
# ---------------------------------------------------------------------------

def _py_record(pid: int) -> dict:
    """One participant-visit with enough T-SICI data to produce a change."""
    r = initialize_record()
    r.update(
        Study="SNBR", ID=pid, Date=VISIT_DATE, Stimulated_cortex="L",
        Subject_type="Patient", Muscle="FDI", Recorded_side="R",
        source_file=f"SNBR-{pid:03d}-TP1C50415A.MEM",
    )
    for isi in TSICI_ISIS:
        r[f"T_SICI_{isi}"] = 55.0
    return r


def _write_redcap_export(path: Path, keys: list[str]):
    """A REDCap data export holding *keys* verbatim, as the real one does."""
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "record_id", "redcap_event_name", "tt_test_date", "cortex",
            "t_sici_p_1_0_ms",
        ])
        for key in keys:
            w.writerow([key, "visit_1_arm_1", VISIT_DATE_ISO, "1", ""])


def _write_dictionary(path: Path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Variable / Field Name", "Form Name"])
        for field in ("t_sici_p_1_0_ms", "tms_values_complete"):
            w.writerow([field, "tms_values"])


def _write_template(path: Path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "record_id", "redcap_event_name", "t_sici_p_1_0_ms",
            "tms_values_complete",
        ])


def _redcap_inputs(directory: Path, keys: list[str]):
    """Write the three REDCap inputs into *directory* and return their paths."""
    directory.mkdir(parents=True, exist_ok=True)
    export = directory / "SNBR_DATA_2026.csv"
    dictionary = directory / "SNBR_DataDictionary_2026.csv"
    template = directory / "SNBR_ImportTemplate_2026.csv"
    _write_redcap_export(export, keys)
    _write_dictionary(dictionary)
    _write_template(template)
    return export, dictionary, template


@pytest.fixture
def redcap_files(tmp_path):
    """The three REDCap inputs, with participant 42 registered as ``042``."""
    return _redcap_inputs(tmp_path / "redcap", keys=["042"])


def _run(py_df, redcap_files, tmp_path, **kwargs):
    export, dictionary, template = redcap_files
    return generate_redcap_import(
        py_df, export, dictionary, template, tmp_path, **kwargs
    )


def _written_ids(output_path: Path) -> list[str]:
    """The record_id column of the emitted CSV, as text off the disk."""
    lines = output_path.read_text(encoding="utf-8-sig").splitlines()
    idx = lines[0].split(",").index("record_id")
    return [line.split(",")[idx] for line in lines[1:] if line]


@pytest.mark.parametrize("pid,expected", [(1, "001"), (10, "010"), (100, "100")])
def test_a_new_participant_is_written_padded(pid, expected, redcap_files, tmp_path):
    py_df = pd.DataFrame([_py_record(pid)])

    import_df, output_path, stats = _run(
        py_df, redcap_files, tmp_path, include_new_ids=True,
    )

    assert stats["new_ids_added"] == [pid]
    assert list(import_df["record_id"]) == [expected]

    # And it reaches the file that way, not just the frame.
    assert _written_ids(output_path) == [expected]


def test_an_existing_participant_is_written_with_redcaps_own_key(
    redcap_files, tmp_path,
):
    """Participant 42 is stored as ``042``; the import must say ``042``."""
    py_df = pd.DataFrame([_py_record(42)])

    import_df, output_path, stats = _run(py_df, redcap_files, tmp_path)

    assert stats["matched"] == 1
    assert stats["new_ids_added"] == []
    assert list(import_df["record_id"]) == ["042"]
    assert _written_ids(output_path) == ["042"]


def test_an_unpadded_redcap_project_keeps_its_unpadded_key(tmp_path):
    """The key is copied, not reformatted -- so ``42`` stays ``42``."""
    files = _redcap_inputs(tmp_path / "unpadded", keys=["42"])
    py_df = pd.DataFrame([_py_record(42)])

    import_df, output_path, _ = _run(py_df, files, tmp_path)

    assert list(import_df["record_id"]) == ["42"]
    assert _written_ids(output_path) == ["42"]


def test_every_emitted_key_names_a_record_redcap_already_holds(tmp_path):
    """The bug as the lab saw it: an unmatched key silently creates a record."""
    keys = ["005", "042", "158"]
    files = _redcap_inputs(tmp_path / "held", keys=keys)
    py_df = pd.DataFrame([_py_record(5), _py_record(42), _py_record(158)])

    import_df, _, stats = _run(py_df, files, tmp_path)

    assert stats["new_ids_added"] == []
    assert len(import_df) == 3
    assert set(import_df["record_id"]) == set(keys)


def test_one_import_can_hold_both_kinds(redcap_files, tmp_path):
    """An existing key and a created one have to survive the same file."""
    py_df = pd.DataFrame([_py_record(42), _py_record(7)])

    import_df, output_path, stats = _run(
        py_df, redcap_files, tmp_path, include_new_ids=True,
    )

    assert stats["new_ids_added"] == [7]
    assert set(import_df["record_id"]) == {"042", "007"}
    assert set(_written_ids(output_path)) == {"042", "007"}


def test_a_skipped_new_participant_writes_nothing_at_all(redcap_files, tmp_path):
    """Without include_new_ids a new participant is dropped, padded or not."""
    py_df = pd.DataFrame([_py_record(7)])

    import_df, _, stats = _run(py_df, redcap_files, tmp_path)

    assert stats["skipped_new_ids"] == [7]
    assert import_df.empty


def test_the_xlsx_change_report_shows_the_same_keys(redcap_files, tmp_path):
    """The report is the human-readable half of the same export.

    Its record_id column reads straight off the emitted rows, so it must show
    what the import will actually touch rather than a different number.
    """
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    py_df = pd.DataFrame([_py_record(42), _py_record(7)])
    reports = tmp_path / "reports"

    _run(
        py_df, redcap_files, tmp_path,
        include_new_ids=True, xlsx_report_dir=str(reports),
    )

    written = sorted(reports.glob("*.xlsx"))
    assert written, "no xlsx change report was written"

    sheet = load_workbook(written[-1]).active
    ids = {
        str(sheet.cell(row, 1).value)
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, 1).value is not None
    }
    assert ids == {"007", "042"}


def test_the_key_is_a_string_so_the_zeros_survive_the_csv(
    redcap_files, tmp_path,
):
    """An int cannot carry a leading zero; the value must stay text.

    Guards the round trip specifically: reading the CSV back with the column
    forced to text must still show the padding, for both kinds of key.
    """
    py_df = pd.DataFrame([_py_record(42), _py_record(1)])

    _, output_path, _ = _run(
        py_df, redcap_files, tmp_path, include_new_ids=True,
    )

    reread = pd.read_csv(output_path, dtype={"record_id": str})
    assert set(reread["record_id"]) == {"042", "001"}


def test_the_per_participant_summary_still_reports_plain_numbers(
    redcap_files, tmp_path,
):
    """The GUI summary counts participants; a padded key must not break it."""
    py_df = pd.DataFrame([_py_record(42), _py_record(7)])

    _, _, stats = _run(
        py_df, redcap_files, tmp_path, include_new_ids=True,
    )

    reported = sorted(e["record_id"] for e in stats["per_participant"])
    assert reported == [7, 42]
    assert all(isinstance(e["record_id"], int) for e in stats["per_participant"])

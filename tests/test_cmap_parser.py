"""Unit tests for the motor nerve conduction study (CMAP) parser.

Two kinds of test live here.

The **synthetic** ones build a PDF and a DOCX in ``tmp_path`` and parse those.
They need pdfplumber / python-docx installed but nothing else, so the parser is
genuinely exercised in any checkout that has the optional dependencies.

The **sample** ones read real lab files from the Y: share. They check the parser
against genuine Viking output, which no synthetic fixture can stand in for, and
skip when the share is not mounted.

Both kinds skip rather than fail when the optional dependency is missing.
``parse_cmap_file`` catches the ImportError, warns, and returns a record with
empty tables -- so without the guard a missing package surfaced as five
assertion errors that read like parser bugs, and
``test_parse_cmap_without_munix_leaves_munix_none`` *passed* for the wrong
reason: MUNIX_table was None because nothing had been parsed at all.

No real patient file is committed here. This repository has two public remotes.
"""

import importlib.util
import json
import sys
from pathlib import Path

import pytest

# Allow imports from the SNBR_TMS_App package root
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from parser.cmap_parser import parse_cmap_file
from processing.df_builder import (
    build_cmap_dataframe,
    build_mem_dataframe,
    merge_cmap_into_mem,
)

# ---------------------------------------------------------------------------
# Optional dependency guards
# ---------------------------------------------------------------------------

requires_pdfplumber = pytest.mark.skipif(
    importlib.util.find_spec("pdfplumber") is None,
    reason="pdfplumber is not installed (optional dependency for PDF CMAP files)",
)
requires_docx = pytest.mark.skipif(
    importlib.util.find_spec("docx") is None,
    reason="python-docx is not installed (optional dependency for DOCX CMAP files)",
)


# ---------------------------------------------------------------------------
# Synthetic fixtures
# ---------------------------------------------------------------------------

CMAP_HEADER = ["Nerve / Site", "Muscle", "Latency", "Amplitude"]
CMAP_UNITS = ["", "", "ms", "mV"]
CMAP_BODY = [
    ["Left Ulnar", "", "", ""],          # group row
    ["Wrist", "FDI", "4.21", "2.2"],
    ["Below Elbow", "FDI", "8.55", "2.0"],
]

MUNIX_HEADER = ["#SIP", "A", "Alpha", "MUNIX", "MUSIX"]
MUNIX_BODY = [["9", "5587", "-0.99", "287", "63"]]


def _write_pdf(path: Path, *, visit_date: str, with_munix: bool) -> Path:
    """A PDF shaped like the EMGRQ reports: a date line and ruled tables.

    Drawn with matplotlib, which the app already depends on. The ruling lines
    matter -- pdfplumber's default table strategy finds tables by their lines,
    so text alone would extract but produce no table.
    """
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    matplotlib.rcParams["pdf.fonttype"] = 42  # embed TrueType so text extracts

    fig = plt.figure(figsize=(8.5, 11))
    fig.text(0.08, 0.95, "Nerve Conduction Study", fontsize=12)
    fig.text(0.08, 0.92, f"Visit Date: {visit_date}", fontsize=10)

    axis = fig.add_axes([0.08, 0.60, 0.84, 0.25])
    axis.axis("off")
    table = axis.table(
        cellText=[CMAP_UNITS] + CMAP_BODY, colLabels=CMAP_HEADER,
        loc="upper left", cellLoc="left",
    )
    table.auto_set_font_size(False)
    table.set_fontsize(9)

    if with_munix:
        munix_axis = fig.add_axes([0.08, 0.30, 0.84, 0.15])
        munix_axis.axis("off")
        munix_table = munix_axis.table(
            cellText=MUNIX_BODY, colLabels=MUNIX_HEADER,
            loc="upper left", cellLoc="left",
        )
        munix_table.auto_set_font_size(False)
        munix_table.set_fontsize(9)

    with PdfPages(path) as pdf:
        pdf.savefig(fig)
    plt.close(fig)
    return path


def _write_docx(path: Path, *, visit_date: str) -> Path:
    """A DOCX shaped like the Natus reports: a date paragraph and a table."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Nerve Conduction Study")
    doc.add_paragraph(f"Visit Date:\t{visit_date}")

    body = [CMAP_HEADER, CMAP_UNITS] + CMAP_BODY
    table = doc.add_table(rows=len(body), cols=len(CMAP_HEADER))
    for r, row in enumerate(body):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    doc.save(str(path))
    return path


@pytest.fixture
def synthetic_pdf(tmp_path):
    return _write_pdf(
        tmp_path / "AA-SNBR-186 AA-SNBR-186.pdf",
        visit_date="14-Apr-26 9:26 AM", with_munix=False,
    )


@pytest.fixture
def synthetic_pdf_with_munix(tmp_path):
    return _write_pdf(
        tmp_path / "AA-SNBR-187 AA-SNBR-187.pdf",
        visit_date="20-Apr-26 10:42 AM", with_munix=True,
    )


@pytest.fixture
def synthetic_docx(tmp_path):
    return _write_docx(
        tmp_path / "AA-SNBR-136 AA-SNBR-136.docx",
        visit_date="3/27/2025 12:27 PM",
    )


# ---------------------------------------------------------------------------
# Synthetic-fixture tests
# ---------------------------------------------------------------------------

@requires_pdfplumber
def test_pdf_yields_participant_date_and_rows(synthetic_pdf):
    record = parse_cmap_file(synthetic_pdf)

    assert record["Study"] == "SNBR"
    assert record["ID"] == 186
    assert record["Date"] == "14/04/2026"
    assert record["source_file"] == synthetic_pdf.name

    rows = json.loads(record["CMAP_table"])
    assert len(rows) == 2, f"expected the two data rows, got {rows}"

    first = rows[0]
    assert first["nerve_site"] == "Left Ulnar / Wrist", (
        "the group row must be folded into the site name, not dropped"
    )
    assert first["muscle"] == "FDI"
    assert first["latency_ms"] == pytest.approx(4.21)
    assert first["amplitude_mv"] == pytest.approx(2.2)


@requires_pdfplumber
def test_pdf_units_row_is_not_read_as_data(synthetic_pdf):
    """A "ms / mV" row would otherwise become a row with no numbers."""
    rows = json.loads(parse_cmap_file(synthetic_pdf)["CMAP_table"])
    assert all(r["latency_ms"] is not None for r in rows)
    assert not any(r["muscle"].lower() in {"ms", "mv"} for r in rows)


@requires_pdfplumber
def test_pdf_without_a_munix_table_leaves_munix_none(synthetic_pdf):
    record = parse_cmap_file(synthetic_pdf)
    assert record["CMAP_table"], "guard: the CMAP table must have parsed"
    assert record["MUNIX_table"] is None


@requires_pdfplumber
def test_pdf_munix_table_is_read_when_present(synthetic_pdf_with_munix):
    record = parse_cmap_file(synthetic_pdf_with_munix)

    assert record["ID"] == 187
    assert record["Date"] == "20/04/2026"
    assert record["MUNIX_table"], "Expected MUNIX_table to be populated"

    rows = json.loads(record["MUNIX_table"])
    assert len(rows) == 1
    row = rows[0]
    assert row["num_sip"] == pytest.approx(9)
    assert row["a"] == pytest.approx(5587)
    assert row["alpha"] == pytest.approx(-0.99)
    assert row["munix"] == pytest.approx(287)
    assert row["musix"] == pytest.approx(63)

    # The MUNIX table must not also be read as CMAP rows.
    cmap_rows = json.loads(record["CMAP_table"])
    assert len(cmap_rows) == 2


@requires_docx
def test_docx_yields_participant_date_and_rows(synthetic_docx):
    record = parse_cmap_file(synthetic_docx)

    assert record["Study"] == "SNBR"
    assert record["ID"] == 136
    assert record["Date"] == "27/03/2025", (
        "the DOCX date is US-order m/d/Y and must not be read as d/m/Y"
    )

    rows = json.loads(record["CMAP_table"])
    assert len(rows) == 2
    assert rows[0]["muscle"] == "FDI"
    assert rows[0]["latency_ms"] == pytest.approx(4.21)


@requires_pdfplumber
def test_merge_writes_the_parsed_table_onto_the_matching_mem_row(synthetic_pdf):
    cmap_df = build_cmap_dataframe([parse_cmap_file(synthetic_pdf)])
    mem_df = build_mem_dataframe([{
        "Study": "SNBR", "ID": 186, "Date": "14/04/2026",
        "Age": 55, "Sex": "M", "Subject_type": "Patient",
        "Stimulated_cortex": "LM", "source_file": "SNBRMC186A.MEM",
    }])

    merged = merge_cmap_into_mem(mem_df, cmap_df)

    assert merged.attrs["cmap_rows_merged"] == 1
    assert merged.attrs["cmap_rows_dropped"] == 0

    row = merged[merged["ID"] == 186].iloc[0]
    assert row["CMAP_table"], "Expected CMAP_table to be populated"
    assert json.loads(row["CMAP_table"])[0]["muscle"] == "FDI"


@requires_pdfplumber
def test_merge_writes_munix_onto_the_matching_mem_row(synthetic_pdf_with_munix):
    cmap_df = build_cmap_dataframe([parse_cmap_file(synthetic_pdf_with_munix)])
    mem_df = build_mem_dataframe([{
        "Study": "SNBR", "ID": 187, "Date": "20/04/2026",
        "Age": 50, "Sex": "F", "Subject_type": "Patient",
        "Stimulated_cortex": "L", "source_file": "SNBRMC187A.MEM",
    }])

    merged = merge_cmap_into_mem(mem_df, cmap_df)

    row = merged[merged["ID"] == 187].iloc[0]
    assert row["MUNIX_table"], "Expected MUNIX_table on the merged row"
    assert json.loads(row["MUNIX_table"])[0]["munix"] == pytest.approx(287)


@requires_pdfplumber
def test_merge_drops_unmatched_cmap_rows(synthetic_pdf):
    """Unmatched CMAP records must be dropped, not appended.

    Appending them would create phantom visit dates with no MEM data that
    the participant panel and Quick Start would then select, blocking every
    other graph type for that visit.
    """
    cmap_df = build_cmap_dataframe([parse_cmap_file(synthetic_pdf)])

    # MEM row has a different ID — no match.
    mem_df = build_mem_dataframe([{
        "Study": "SNBR", "ID": 999, "Date": "01/01/2020",
        "Age": 60, "Sex": "F", "Subject_type": "Control",
        "Stimulated_cortex": "LM", "source_file": "SNBRMC999A.MEM",
    }])

    merged = merge_cmap_into_mem(mem_df, cmap_df)

    assert merged.attrs["cmap_rows_merged"] == 0
    assert merged.attrs["cmap_rows_dropped"] == 1
    # Only the original MEM row — no phantom CMAP row appended.
    assert len(merged) == 1
    assert int(merged["ID"].iloc[0]) == 999


def test_an_unsupported_extension_is_reported_not_crashed(tmp_path):
    """A stray file in the CMAP folder must not take the import down."""
    stray = tmp_path / "AA-SNBR-186 notes.txt"
    stray.write_text("not a nerve conduction study", encoding="utf-8")

    with pytest.warns(UserWarning, match="Unsupported CMAP file extension"):
        record = parse_cmap_file(stray)

    assert record["ID"] == 186
    assert record["CMAP_table"] is None


# ---------------------------------------------------------------------------
# Sample files from the lab share
# ---------------------------------------------------------------------------
# Real Viking output, which no synthetic fixture reproduces: the true column
# spellings, merged cells and date formats. Skipped when the share is absent.

_SAMPLE_PDF = Path(
    r"Y:\Merged Data\Viking\EMGRQ\AA-SNBR-186 AA-SNBR-186, AA-SNBR-186   14-Apr-26 9-26 AM.pdf"
)
_SAMPLE_DOCX = Path(
    r"Y:\Merged Data\Viking\Industry\AA-SNBR-136 AA-SNBR- (AA-SNBR-136) 3_27_2025 12_27_44 PM R1.docx"
)
_SAMPLE_PDF_MUNIX = Path(
    r"Y:\Merged Data\Viking\EMGRQ\AA-SNBR-187 AA-SNBR-187, AA-SNBR-187   20-Apr-26 10-42 AM.pdf"
)

requires_samples = pytest.mark.skipif(
    not (_SAMPLE_PDF.exists() and _SAMPLE_DOCX.exists()),
    reason="Sample CMAP files on the Y: share are not available.",
)


@requires_samples
@requires_pdfplumber
def test_sample_pdf_extracts_participant_date_and_rows():
    record = parse_cmap_file(_SAMPLE_PDF)

    assert record["Study"] == "SNBR"
    assert record["ID"] == 186
    assert record["Date"] == "14/04/2026"
    assert record["source_file"] == _SAMPLE_PDF.name

    rows = json.loads(record["CMAP_table"])
    assert len(rows) == 2

    first = rows[0]
    assert "Ulnar" in first["nerve_site"]
    assert first["muscle"] == "FDI"
    assert first["latency_ms"] == pytest.approx(4.21)
    assert first["amplitude_mv"] == pytest.approx(2.2)


@requires_samples
@requires_docx
def test_sample_docx_extracts_participant_date_and_rows():
    record = parse_cmap_file(_SAMPLE_DOCX)

    assert record["Study"] == "SNBR"
    assert record["ID"] == 136
    assert record["Date"] == "27/03/2025"

    rows = json.loads(record["CMAP_table"])
    assert len(rows) >= 1
    first = rows[0]
    assert "Median" in first["nerve_site"]
    # The docx file shows Median APB on the Wrist site
    assert first["muscle"] == "APB"
    assert first["latency_ms"] == pytest.approx(2.77)
    assert first["amplitude_mv"] == pytest.approx(14.3)


@requires_samples
@requires_pdfplumber
def test_sample_pdf_extracts_munix_when_present():
    if not _SAMPLE_PDF_MUNIX.exists():
        pytest.skip("SNBR-187 sample PDF with MUNIX not available.")
    record = parse_cmap_file(_SAMPLE_PDF_MUNIX)

    assert record["ID"] == 187
    assert record["Date"] == "20/04/2026"
    assert record["MUNIX_table"], "Expected MUNIX_table to be populated"

    rows = json.loads(record["MUNIX_table"])
    assert len(rows) == 1
    row = rows[0]
    assert row["num_sip"] == pytest.approx(9)
    assert row["a"] == pytest.approx(5587)
    assert row["alpha"] == pytest.approx(-0.99)
    assert row["munix"] == pytest.approx(287)
    assert row["musix"] == pytest.approx(63)


@requires_samples
@requires_pdfplumber
def test_sample_merge_writes_onto_matching_mem_row():
    record = parse_cmap_file(_SAMPLE_PDF)
    cmap_df = build_cmap_dataframe([record])

    mem_df = build_mem_dataframe([{
        "Study": "SNBR", "ID": 186, "Date": "14/04/2026",
        "Age": 55, "Sex": "M", "Subject_type": "Patient",
        "Stimulated_cortex": "LM", "source_file": "SNBRMC186A.MEM",
    }])

    merged = merge_cmap_into_mem(mem_df, cmap_df)

    assert merged.attrs["cmap_rows_merged"] == 1
    assert merged.attrs["cmap_rows_dropped"] == 0

    row = merged[merged["ID"] == 186].iloc[0]
    assert row["CMAP_table"], "Expected CMAP_table to be populated"
    parsed = json.loads(row["CMAP_table"])
    assert len(parsed) == 2
    assert parsed[0]["muscle"] == "FDI"

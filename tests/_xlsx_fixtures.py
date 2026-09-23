"""Synthetic QtracP workbooks for the xlsx tests.

Real exports are patient data on two public remotes and must never be
committed, so the tests build their own. The layout mirrors QtracP's: one sheet
per variable (``D``, ``P``, ``L``, ``T``), and inside each a ``(time, value)``
column pair per channel under a ``Chan  N`` header in the value column.

The recording built by :func:`tsici_recording` is arranged so that Qtrac's
weighted log regression returns its thresholds *exactly*: every MEP follows
``0.2 mV * exp(k * (T - threshold))``, so the regression of T on ln(P) is a
perfect fit and the weights cannot move it.
"""

from __future__ import annotations

import math
from pathlib import Path

TOKEN = "TP3C60520A"
SLOPE = 0.5                       # ln(MEP) per %MSO around threshold
REF_THRESHOLD = 40.0              # channel-1 RMT200, %MSO
TSICI_THRESHOLDS = {1.0: 44.0, 1.5: 36.0}   # -> 110 % and 90 % of the reference
ASICF_VALUES = {1.0: 400.0, 1.3: 100.0}     # -> as % of a baseline of exactly 1.0 mV


def mep_for(stimulus: float, threshold: float) -> float:
    return 0.2 * math.exp(SLOPE * (stimulus - threshold))


def write_workbook(path: Path, channels: dict[int, list[tuple]], sheets=("D", "P", "L", "T")) -> Path:
    """Write ``channels`` ({chan: [(time, T, P, L, D), ...]}) in QtracP's layout."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    index = {"T": 1, "P": 2, "L": 3, "D": 4}
    order = sorted(channels)
    for sheet in sheets:
        ws = wb.create_sheet(sheet)
        if sheet not in index:                       # a waveform sheet, e.g. "M1-1"
            ws.append([None, "Chan 1: Trace 1"])
            for i in range(20):
                ws.append([i / 10.0, 0.001 * i])
            continue
        header = [None]
        for _chan in order:
            header += [None, None, None]
        for k, chan in enumerate(order):
            header[1 + 3 * k + 1] = f"Chan  {chan}"
        # openpyxl writes None cells as empty, so the header row is sparse as in the real files
        ws.append(header[: 1 + 3 * len(order)])
        n_rows = max(len(v) for v in channels.values())
        for r in range(n_rows):
            row = [None] * (1 + 3 * len(order))
            for k, chan in enumerate(order):
                pulses = channels[chan]
                if r < len(pulses):
                    time, T, P, L, D = pulses[r]
                    row[1 + 3 * k] = time
                    row[1 + 3 * k + 1] = (T, P, L, D)[index[sheet] - 1]
            ws.append(row)
    ws = wb.create_sheet("QPP")
    ws.append([None])
    wb.save(path)
    return path


def write_waveform_workbook(path: Path) -> Path:
    """A CSP export: waveform sheets only, nothing per stimulus."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in ("Wxxxx", "Wxxx", "Wxx", "Wx", "W"):
        ws = wb.create_sheet(name)
        ws.append([None, "Chan 1: Trace 5"])
        for i in range(10):
            ws.append([i / 10.0, 0.01 * i])
    wb.create_sheet("QPP").append([None])
    wb.save(path)
    return path


def tsici_recording(*, increment_mode: bool = False) -> dict[int, list[tuple]]:
    """A T-SICI phase (5 to 9 min) followed by an A-SICF phase (15 to 17 min).

    Channel 1 tracks RMT200 throughout at a true threshold of 40 %MSO. Channel
    11 (ISI 1.0 ms) has a true threshold of 44, channel 12 (1.5 ms) of 36. Three
    EMG-rejected traces sit on channel 20 inside the phase. In the A-SICF phase
    the pairs on channels 11 (-1.0 ms) and 12 (-1.3 ms) give geometric means of
    4.0 and 1.0 mV against a baseline of exactly 1.0 mV: the channel-5 pulses
    at 14.9, 15.5 and 16.5 min (the one just before the first pair, then those
    up to the last pair at 17.0). The 8 mV pulse at 12.0 and the 1 mV pulse at
    17.5 fall outside that window and must be left out.
    """
    chans: dict[int, list[tuple]] = {}
    # channel 1: tracking pulses before and during the T-SICI phase
    c1 = []
    for i, T in enumerate([36, 38, 40, 42, 44, 40, 39, 41, 40, 42, 38, 40, 41, 39]):
        c1.append((1.0 + i * 0.6, T, mep_for(T, REF_THRESHOLD), 25.0, 0))
    chans[1] = c1

    def latest_c1(t: float) -> float:
        """The channel-1 stimulus most recently delivered before *t* (what the
        conditioning pulse is set from, so what an increment is relative to)."""
        return [p[1] for p in c1 if p[0] <= t][-1]

    # conditioned channels, tracking around their own thresholds
    for chan, isi, stimuli in (
        (11, 1.0, [42, 43, 44, 45, 46, 44]),
        (12, 1.5, [34, 35, 36, 37, 38, 36]),
    ):
        pulses = []
        for i, T in enumerate(stimuli):
            time = 5.0 + i * 0.7 + (0.1 if chan == 12 else 0.0)
            P = mep_for(T, TSICI_THRESHOLDS[isi])
            exported_T = T - round(0.7 * latest_c1(time)) if increment_mode else T
            pulses.append((time, exported_T, P, 26.0, isi))
        chans[chan] = pulses
    # rejected traces: absurd stimuli that must never be counted
    chans[20] = [(5.2, 99, 5.0, 20.0, 1.0), (6.4, 99, 5.0, 20.0, 1.0), (7.1, 99, 5.0, 20.0, 1.5)]
    # A-SICF phase: channel 5 is the test-alone pulse, pairs at a fixed 50 %MSO
    chans[5] = [
        (12.0, 50, 8.0, 24.0, 0),        # before the phase: excluded from the baseline
        (14.9, 50, 1.0, 24.0, 0),        # the pulse just before the first pair: included
        (15.5, 50, 2.0, 24.0, 0),
        (16.5, 50, 0.5, 24.0, 0),
        (17.5, 50, 1.0, 24.0, 0),
    ]
    chans[11] += [(15.1, 50, 2.0, 24.0, -1.0), (15.9, 50, 4.0, 24.0, -1.0), (16.9, 50, 8.0, 24.0, -1.0)]
    chans[12] += [(15.2, 50, 1.0, 24.0, -1.3), (16.0, 50, 1.0, 24.0, -1.3), (17.0, 50, 1.0, 24.0, -1.3)]
    return chans


def write_tsici_workbook(folder: Path, name: str = f"SNBR-192-{TOKEN}.xlsx", **kwargs) -> Path:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    return write_workbook(folder / name, tsici_recording(**kwargs))
